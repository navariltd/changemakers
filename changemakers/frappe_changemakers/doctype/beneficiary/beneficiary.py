# Copyright (c) 2022, hussain@frappe.io and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.contacts.address_and_contact import load_address_and_contact
from frappe.model.document import Document
from frappe.query_builder import Field
from frappe.utils import cint

import csv
import io
from frappe.utils.file_manager import get_file
from openpyxl import load_workbook


from changemakers.utils.data import is_valid_indian_phone_number

class Beneficiary(Document):
    def before_insert(self):
        settings = frappe.get_doc("Changemakers Settings")
        if settings.enable_student_creation_on_beneficiary and self.beneficiary_type == "Student":
            # Create Student
            student = frappe.new_doc("Student")
            student.first_name = self.first_name
            student.student_email_id = self.email
            student.last_name = self.last_name
            student.beneficiary = self.name  
            student.insert(ignore_permissions=True)
            self.student = student.name
            
    def after_insert(self):
        # delete user and customer assocaited with beneficiary if exists
        settings = frappe.get_doc("Changemakers Settings")
        if settings.enable_student_creation_on_beneficiary:
            student_name = frappe.db.get_value("Student", {"student_email_id": self.email})
            if not student_name:
                return
            
            student = frappe.get_doc("Student", student_name)
            user = frappe.db.get_value("User", {"email": self.email})
            customer = frappe.db.get_value("Customer", {"name": student.customer})
            frappe.db.set_value("Student", student_name, {"user": None, "customer": None})
            
            if customer:
                frappe.delete_doc("Customer", customer, ignore_permissions=True)
            if user and user != frappe.session.user:  
                frappe.delete_doc("User", user, ignore_permissions=True)

        if settings.generate_supplier_when_beneficiary_is_created:
            self.create_supplier()

    def on_trash(self):
        if self.supplier:
            frappe.db.set_value("Supplier", self.supplier, "disabled", 1)
        
    def before_save(self):
        self.set_created_by()
        self.full_name = f"{self.first_name} {self.last_name or ''}"

    def validate(self):
        self.validate_age()
        settings = frappe.get_doc("Changemakers Settings")
        if not self.status and getattr(settings, "default_beneficiary_status", None):
            self.status = settings.default_beneficiary_status
        # Check if manage_beneficiary_lifecycle is enabled
        if settings.manage_beneficiary_lifecycle:
            if self.status == "Active" and self.recruitment_phase and not self.beneficiary_no:
                self.validate_available_slots()
                self.beneficiary_no = generate_beneficiary_no(self)
            if self.status in ["Disqualified", "Relocated"] and not self.archive_date:
                self.archive_date = frappe.utils.nowdate()

            if self.status == "Active" and not self.activation_date:
                self.activation_date = frappe.utils.nowdate()

        
    def validate_available_slots(self):
        if not self.recruitment_phase:
            return
        recruitment_phase = frappe.get_doc("Recruitment Phase", self.recruitment_phase)

        available_slots = recruitment_phase.available_slots
        if available_slots <= 0:
            frappe.throw(
                _(f"No available slots in {self.recruitment_phase}.")
            )

    def validate_age(self):
        if self.age and not (self.age < 120):
            frappe.throw(f"Value of {frappe.bold('Age')} should be less than 120!")

    def validate_phone_number_fields(self):
        if self.poc_phone and not is_valid_indian_phone_number(self.poc_phone):
            frappe.throw(
                f"Value of {frappe.bold('POC Phone')} is not a valid Indian Phone number"
            )

        if self.phone_number and not is_valid_indian_phone_number(self.phone_number):
            frappe.throw(
                f"Value of {frappe.bold('Phone')} is not a valid Indian Phone number"
            )

    def set_created_by(self):
        if not self.created_by:
            owner = frappe.db.get_value("User", self.owner, "full_name")
            self.created_by = owner

    def onload(self):
        load_address_and_contact(self)

    @frappe.whitelist()
    def create_supplier(self):
        if self.supplier:
            return
        
        existing_supplier = frappe.db.get_value("Supplier", {"supplier_name": self.full_name})
        if existing_supplier:
            self.supplier = existing_supplier
            self.save()
            return existing_supplier
        
        if not frappe.db.exists("Supplier Group", "Beneficiary"):
            supplier_group = frappe.new_doc("Supplier Group")
            supplier_group.supplier_group_name = "Beneficiary"
            supplier_group.flags.ignore_permissions = True
            supplier_group.insert()

        supplier = frappe.new_doc("Supplier")
        supplier.supplier_name = self.full_name
        supplier.supplier_type = "Individual"
        supplier.supplier_group = "Beneficiary"

        supplier.flags.ignore_permissions = True
        supplier.insert()
        self.supplier = supplier.name
        self.save()
        return supplier.name


def generate_beneficiary_no(doc):
    recruitment_phase = frappe.get_doc("Recruitment Phase", doc.recruitment_phase)
    
    Beneficiary = frappe.qb.DocType("Beneficiary")
    beneficiary_no_field = Field("beneficiary_no")
    
    max_number_query = (
        frappe.qb.from_(Beneficiary)
        .select(beneficiary_no_field)
        .where(
            (Beneficiary.recruitment_phase == doc.recruitment_phase) &
            (Beneficiary.beneficiary_no.like(f"{recruitment_phase.beneficiary_number_series}%"))
        )
    )

    beneficiary_nos = frappe.db.sql(max_number_query, as_dict=False)

    numeric_parts = [
        cint(bn[0][-4:])  
        for bn in beneficiary_nos if bn[0][-4:].isdigit()
    ]

    last_number = max(numeric_parts, default=0)

    new_number = last_number + 1
    formatted_number = f"{new_number:04}"

    beneficiary_no = f"{recruitment_phase.beneficiary_number_series}{formatted_number}"

    return beneficiary_no


@frappe.whitelist()
def upload_beneficiary_list(file_url):
    doctype = "Beneficiary"
    file_doc = get_file(file_url)
    content = file_doc[1]
    filename = file_doc[0]

    errors = []
    processed = []
    rows = []

    if filename.lower().endswith(".csv"):
        try:
            stream = io.StringIO(content.decode("utf-8"))
            reader = csv.DictReader(stream)
            rows = list(reader)
        except Exception as e:
            frappe.throw(f"Failed to read CSV: {str(e)}")

    elif filename.lower().endswith((".xlsx", ".xls")):
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            headers = [str(cell.value).strip() for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append(dict(zip(headers, row)))
        except Exception as e:
            frappe.throw(f"Failed to read Excel: {str(e)}")
    else:
        frappe.throw("Unsupported file type. Upload CSV or Excel.")

    meta = frappe.get_meta(doctype)
    valid_fields = {df.fieldname for df in meta.fields}
    donor_field = "donor"
    beneficiary_no_field = "beneficiary_no"
    bank_fields = {"bank_name", "bank_branch_name", "bank_account_number", "account_holder_name"}

    for idx, row in enumerate(rows, start=2):
        if not any(str(v).strip() if v is not None else "" for v in row.values()):
            continue

        beneficiary = None
        
        try:
            if row.get("name") and frappe.db.exists(doctype, row.get("name")):
                beneficiary = frappe.get_doc(doctype, row.get("name"))
            else:
                beneficiary = frappe.new_doc(doctype)
        except Exception as e:
            errors.append({"row": idx, "error": f"Doc Initialization Error: {str(e)}"})
            continue

        for field, value in row.items():
            if field in (donor_field, beneficiary_no_field) or field in bank_fields or field == "name" or value in [None, "None"]:
                continue
            
            if field in valid_fields:
                try:
                    beneficiary.set(field, value)
                except Exception as e:
                    errors.append({"row": idx, "field": field, "error": str(e)})
            else:
                errors.append({"row": idx, "field": field, "error": f"Field '{field}' does not exist"})

        try:
            beneficiary.save(ignore_permissions=True)
        except Exception as e:
            errors.append({"row": idx, "error": f"Initial Save Failed: {str(e)}"})
            continue

        donor_name = row.get(donor_field)
        beneficiary_no = row.get(beneficiary_no_field)
        if donor_name and beneficiary_no:
            try:
                donor_ref = donor_name
                if not frappe.db.exists("Donor", donor_name):
                    donor_doc = frappe.get_doc("Donor", {"donor_name": donor_name})
                    donor_ref = donor_doc.name
                
                exists = False
                for d in beneficiary.get("donors") or []:
                    if d.donor == donor_ref:
                        d.beneficiary_no = beneficiary_no
                        exists = True
                        break
                if not exists:
                    beneficiary.append("donors", {"donor": donor_ref, "beneficiary_no": beneficiary_no})
                beneficiary.save(ignore_permissions=True)
            except Exception as e:
                errors.append({"row": idx, "field": "donors_child_table", "error": str(e)})
        elif donor_name or beneficiary_no:
            errors.append({"row": idx, "error": "Both donor and beneficiary_no must be provided together"})

        bank_data = {f: row.get(f) for f in bank_fields if row.get(f)}
        if bank_data:
            try:
                if len(bank_data) != len(bank_fields):
                    missing = bank_fields - bank_data.keys()
                    errors.append({"row": idx, "error": f"Incomplete bank details. Missing: {', '.join(missing)}"})
                else:
                    bank = frappe.db.exists({
                        "doctype": "Bank",
                        "bank_name": bank_data["bank_name"],
                    })
                    if bank:
                        bank_doc = frappe.get_doc("Bank", {"bank_name": bank_data["bank_name"]})
                    else:
                        bank_doc = frappe.new_doc("Bank")
                        bank_doc.bank_name = bank_data["bank_name"]
                        bank_doc.flags.ignore_permissions = True
                        bank_doc.insert()

                    existing_account = frappe.db.exists({
                        "doctype": "Bank Account",
                        "bank": bank_doc.name,
                        "bank_account_no": bank_data["bank_account_number"],
                    })
                    
                    if not existing_account:
                        bank_account = frappe.new_doc("Bank Account")
                        bank_account.bank = bank_doc.name
                        bank_account.bank_account_no = bank_data["bank_account_number"]
                        bank_account.branch_code = bank_data["bank_branch_name"]
                        bank_account.account_name = bank_data["account_holder_name"]
                        bank_account.party_type = "Supplier"
                        bank_account.party = beneficiary.supplier
                        bank_account.flags.ignore_permissions = True
                        bank_account.insert()
                    
                    beneficiary.save(ignore_permissions=True)
            except Exception as e:
                errors.append({"row": idx, "field": "bank_details", "error": str(e)})

        processed.append(beneficiary.name)

    return {
        "processed": processed,
        "errors": errors,
    }